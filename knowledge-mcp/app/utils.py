import os
import hashlib
import logging
from pathlib import Path
from typing import List, Dict, Any

import yaml
import requests

logger = logging.getLogger(__name__)

ALLOWED_EXT = {".md", ".txt", ".pdf", ".docx", ".doc"}


def load_config(config_path: str = "/app/config.yaml") -> Dict[str, Any]:
    with open(config_path, "r", encoding="utf-8") as f:
        raw = f.read()
    for key, val in [
        ("QDRANT_HOST", os.getenv("QDRANT_HOST", "qdrant")),
        ("QDRANT_PORT", os.getenv("QDRANT_PORT", "6333")),
        ("MODEL_API_BASE", os.getenv("MODEL_API_BASE", "https://api.openai.com/v1")),
        ("EMBED_MODEL", os.getenv("EMBED_MODEL", "text-embedding-3-large")),
        ("RERANK_MODEL", os.getenv("RERANK_MODEL", "rerank-1")),
        ("KB_DATA_DIR", os.getenv("KB_DATA_DIR", "/data/mcpdata")),
    ]:
        raw = raw.replace(f"${{{key}}}", val)
    return yaml.safe_load(raw)


def get_env(name: str, default: str = "") -> str:
    return os.getenv(name, default)


def read_text_file(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        logger.warning("读取文本文件失败 %s: %s", path, e)
        return ""


def read_pdf(path: Path) -> str:
    try:
        import fitz
        doc = fitz.open(str(path))
        texts = []
        for page in doc:
            texts.append(page.get_text())
        doc.close()
        return "\n".join(texts)
    except ImportError:
        logger.warning("pymupdf 未安装，无法读取 PDF: %s", path)
        return ""
    except Exception as e:
        logger.warning("读取 PDF 失败 %s: %s", path, e)
        return ""


def read_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        texts = []
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        texts.append(cell.text)
        return "\n".join(texts)
    except ImportError:
        logger.warning("python-docx 未安装，无法读取 DOCX: %s", path)
        return ""
    except Exception as e:
        logger.warning("读取 DOCX 失败 %s: %s", path, e)
        return ""


def read_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return read_text_file(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext in (".docx", ".doc"):
        return read_docx(path)
    return ""


def chunk_text(text: str, chunk_size: int, overlap: int) -> List[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    text_len = len(text)
    while start < text_len:
        end = min(start + chunk_size, text_len)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= text_len:
            break
        start = max(end - overlap, start + 1)
    return chunks


def make_point_id(relative_path: str, chunk_index: int) -> str:
    raw = f"{relative_path}::{chunk_index}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()


def file_hash(path: Path) -> str:
    h = hashlib.md5()
    h.update(str(path.stat().st_mtime_ns).encode())
    h.update(str(path.stat().st_size).encode())
    return h.hexdigest()


def get_embeddings(texts: List[str], cfg: Dict[str, Any]) -> List[List[float]]:
    api_base = cfg["models"]["api_base"]
    api_key = get_env("MODEL_API_KEY", "")
    model = cfg["models"]["embed_model"]
    if not api_key:
        raise ValueError("Missing required environment variable MODEL_API_KEY")
    url = f"{api_base}/embeddings"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    all_embeddings = []
    batch_size = 64
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        payload = {"model": model, "input": batch}
        resp = requests.post(url, headers=headers, json=payload, timeout=120)
        resp.raise_for_status()
        data = resp.json()
        all_embeddings.extend(item["embedding"] for item in data["data"])
    return all_embeddings


def rerank(query: str, documents: List[str], top_n: int, cfg: Dict[str, Any]) -> List[Dict[str, Any]]:
    if not documents:
        return []
    api_base = cfg["models"]["api_base"]
    api_key = get_env("MODEL_API_KEY", "")
    model = cfg["models"]["rerank_model"]
    if not api_key:
        raise ValueError("Missing required environment variable MODEL_API_KEY")
    url = f"{api_base}/rerank"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": model,
        "query": query,
        "documents": documents,
        "top_n": top_n,
    }
    resp = requests.post(url, headers=headers, json=payload, timeout=120)
    resp.raise_for_status()
    return resp.json().get("results", [])
